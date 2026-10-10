#!/usr/bin/env python3
"""Build the Review-1 panel deck into the official PPTX template.

Source of truth : Review1_audio_trng.md  (sections map 1:1 to the template)
Template        : ../1. IDP Review 1 template.pptx   (blank copy kept alongside this md)
Output          : ../1. IDP Review 1 template.pptx   (filled in place, written atomically)
Diagrams        : ../diagrams/fig*.png  (fig5 excluded on purpose - see notes)

Re-run after editing the markdown or filling the team/guide placeholders:
    cd ~/audio-trng && uv run python "idp review 1/build_review1_pptx.py"
"""

from __future__ import annotations

import copy
import os
import re
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TPL = os.path.join(ROOT, "1. IDP Review 1 template.pptx")
MD = os.path.join(HERE, "Review1_audio_trng.md")
DIAG = os.path.join(ROOT, "diagrams")

# palette taken from the template itself
NAVY = "14285E"  # headings / labels
DARK = "1F2937"  # body copy
GRAY = "5B6472"  # footer
WHITE = "FFFFFF"
HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

IN = Inches


# --------------------------------------------------------------------------- text


def clean(s: str) -> str:
    """Strip markdown emphasis/code markers - PPTX renders nothing of ** or `."""
    s = s.replace("\u00a0", " ").replace("\u2011", "-")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"(?<!\S)\*(?!\s)(.+?)(?<!\s)\*(?!\S)", r"\1", s)
    return re.sub(r"\s+", " ", s).strip()


def trunc(s: str, n: int) -> str:
    """Clean then hard-truncate at a word boundary with an ellipsis."""
    s = clean(s)
    if len(s) <= n:
        return s
    cut = s[: max(0, n - 1)]
    return (cut.rsplit(" ", 1)[0] if " " in cut else cut) + "\u2026"


def fill(shape, paragraphs, *, anchor=None, margins=None):
    """Rewrite a text frame with explicit run formatting.

    paragraphs: list of list of run dicts {t, sz, b, i, c, font}
    """
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    if margins is not None:
        tf.margin_left, tf.margin_right = margins[0], margins[1]
        tf.margin_top, tf.margin_bottom = margins[2], margins[3]
    tf.clear()
    for pi, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = runs[0].get("align") if runs else None
        if runs and runs[0].get("space_after") is not None:
            p.space_after = Pt(runs[0]["space_after"])
        if runs and runs[0].get("line") is not None:
            p.line_spacing = runs[0]["line"]
        for spec in runs:
            r = p.add_run()
            r.text = spec["t"]
            f = r.font
            f.size = Pt(spec.get("sz", 12))
            f.bold = spec.get("b", False)
            f.italic = spec.get("i", False)
            f.name = spec.get("font", BODY_FONT)
            f.color.rgb = RGBColor.from_string(spec.get("c", DARK))
    return shape


def by_name(slide, name, kind=None, nth=0):
    hits = [
        s
        for s in slide.shapes
        if s.name == name and (kind is None or s.shape_type == kind)
    ]
    if not hits:
        raise KeyError(f"shape {name!r} not on slide")
    return hits[nth]


def dup_slide(prs, idx):
    """Clone a slide (shapes only - safe for the table/text-only template slides)."""
    src = prs.slides[idx]
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):  # drop placeholders inherited from the layout
        dst.shapes._spTree.remove(sh._element)
    for sh in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst


# --------------------------------------------------------------------- markdown


def md_lines():
    return open(MD, encoding="utf-8").read().split("\n")


def md_section(lines, start, end):
    si = next(i for i, l in enumerate(lines) if re.match(start, l))
    ei = len(lines)
    if end:
        for i in range(si + 1, len(lines)):
            if re.match(end, lines[i]):
                ei = i
                break
    return lines[si:ei]


def md_table(sec):
    rows, cur = [], None
    for l in sec:
        if l.startswith("|"):
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            rows.append(cells)
        elif rows:
            break
    return rows



def parse_content():
    L = md_lines()
    out = {}

    # --- domain / application area (slide 2 header)
    brief = md_section(L, r"## 1\. Brief", r"### Problem")
    dom = next(l for l in brief if l.startswith("**Domain:**"))
    app = next(l for l in brief if l.startswith("**Application Area:**"))
    out["domain"] = clean(dom.split("**Domain:**", 1)[1])
    out["app"] = clean(app.split("**Application Area:**", 1)[1])

    # --- 3-column problem / solution / outcome table
    tri = md_table(md_section(L, r"\| \*\*1\. Problem", r"\*\*Why the approach"))
    body = tri[1]
    out["tri"] = [clean(c) for c in body]

    # --- 6 key techniques
    tech = md_table(md_section(L, r"## 2\. Key", r"## 3\. Literature"))
    out["tech"] = [(clean(r[1]), clean(r[2])) for r in tech[1:]]

    # --- literature review, 3 themes
    themes = []
    for a, b in (
        (r"### 3\.1", r"### 3\.2"),
        (r"### 3\.2", r"### 3\.3"),
        (r"### 3\.3", r"\*\*Reading the table"),
    ):
        rows = md_table(md_section(L, a, b))
        themes.append([[clean(c) for c in r] for r in rows[1:]])
    out["themes"] = themes

    # --- 6 research gaps: title + the italic "*Gap: ...*" payload
    gs = md_section(L, r"## 4\. Research Gaps", r"## 5\. Objectives")
    gaps = []
    for i, l in enumerate(gs):
        m = re.match(r"\*\*(G\d+) — (.+?)\*\*", l)
        if not m:
            continue
        j, body = i + 1, []
        while j < len(gs) and not re.match(r"\*\*G\d+ —", gs[j]):
            if gs[j].strip():
                body.append(gs[j].strip())
            j += 1
        blob = " ".join(body)
        gm = re.search(r"\*Gap:\s*(.+?)\*\s*$", blob, re.S)
        gaps.append(
            {
                "id": m.group(1),
                "title": clean(m.group(2).rstrip(".")),
                "gap": clean(gm.group(1)) if gm else clean(blob),
            }
        )
    out["gaps"] = gaps

    # --- main objective + O1..O4
    obj_sec = md_section(L, r"## 5\. Objectives", r"## 6\. References")
    main = next(l for l in obj_sec if l.startswith("To design, implement"))
    out["main_obj"] = clean(main)
    orows = md_table(obj_sec)
    objs = []
    for r in orows[1:]:
        objs.append(
            {
                "id": clean(r[0]),
                "title": clean(r[1]).split(":", 1)[0],
                "detail": clean(r[1]).split(":", 1)[1].strip()
                if ":" in r[1]
                else "",
                "target": clean(r[2]),
                "gap": clean(r[3]),
            }
        )
    out["objs"] = objs

    # --- references: numeric [1]-[72] then standards [S1]-[S5]
    rs = md_section(L, r"## 6\. References", r"## 7\. Appendix")
    refs = []
    for l in rs:
        m = re.match(r"(\[\d+\]|\[S\d+\])\s+(.*)", l.strip())
        if m:
            refs.append((m.group(1), clean(m.group(2))))
    out["refs"] = refs
    return out



# ------------------------------------------------------------------------- layout


def fit_size(text, w_in, h_in, hi, lo, *, safety=0.95, lines_max=None):
    """Pick the largest font size (pt) whose wrapped text fits the box."""
    for sz in [x / 2 for x in range(int(hi * 2), int(lo * 2) - 1, -1)]:
        cpl = (w_in - 0.2) * 72 / (0.48 * sz)
        lines = (h_in - 0.1) * 72 / (sz * 1.22)
        if lines_max:
            lines = min(lines, lines_max)
        if len(text) <= cpl * lines * safety:
            return sz
    return lo


def set_cell(cell, text, sz, *, bold=False, color=DARK, font=BODY_FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = RGBColor.from_string(color)


def add_row(table):
    table._tbl.append(copy.deepcopy(table._tbl.tr_lst[-1]))


def clone_shapes(slide, shapes):
    out = []
    for sh in shapes:
        slide.shapes._spTree.append(copy.deepcopy(sh._element))
        out.append(slide.shapes[-1])
    return out


# ------------------------------------------------------------------------ slide 1


def build_title(s, C):
    fill(by_name(s, "Text 2"),
         [[{"t": "PROJECT ID:  [PROJECT ID]", "sz": 20, "b": True, "c": NAVY}]])
    fill(by_name(s, "Text 3"),
         [[{"t": "Generation of True Random Numbers from Bird Chirping",
            "sz": 34, "b": True, "c": NAVY, "font": HEAD_FONT}]])

    table = next(sh.table for sh in s.shapes if sh.has_table)
    while len(table.rows) > 5:
        table._tbl.remove(table._tbl.tr_lst[-1])
    for i in range(4):
        set_cell(table.cell(i + 1, 1), f"[Name {i + 1}]", 11)
        set_cell(table.cell(i + 1, 2), f"[Reg No {i + 1}]", 11)
        set_cell(table.cell(i + 1, 3), "SCOPE", 11)

    fill(by_name(s, "Text 7"),
         [
             [{"t": "Guide Name", "sz": 20, "b": True, "c": NAVY}],
             [{"t": "Designation", "sz": 14, "c": DARK}],
             [{"t": "School: SCOPE", "sz": 14, "c": DARK}],
         ])



# --------------------------------------------------------------------- slide 4+


LIT_COLS = [8, 50, 12, 63, 50, 70, 67]  # per-column char budgets at 6.5pt / 2 lines


def build_lit(s, rows, title, kicker, note):
    """One literature-review slide: up to 11 paper rows under a theme heading."""
    sh_title = by_name(s, "Text 0")
    sh_title.width = Inches(12.13)
    fill(sh_title,
         [[{"t": title, "sz": 34, "b": True, "c": NAVY, "font": HEAD_FONT}]])

    fill(by_name(s, "Text 4", nth=1),               # "Add as many slides..."
         [[{"t": kicker, "sz": 12, "b": True, "c": NAVY}]])
    fill(by_name(s, "Text 4"),                      # guidance line at the bottom
         [[{"t": note, "sz": 9, "i": True, "c": GRAY}]])

    gf = next(sh for sh in s.shapes if sh.has_table)
    table = gf.table
    while len(table.rows) < len(rows) + 1:
        add_row(table)
    table.rows[0].height = Inches(0.42)
    for i in range(1, len(rows) + 1):
        table.rows[i].height = Inches(0.36)
    gf.height = Inches(0.42 + len(rows) * 0.36)

    set_cell(table.cell(0, 0), "Ref.", 7.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    for c, h in enumerate(
        ["Author(s), Year", "Category / Theme", "Proposed Technique",
         "Dataset / Tools", "Key Findings", "Limitations"], start=1):
        set_cell(table.cell(0, c), h, 7.5, bold=True, color=WHITE)

    theme = re.search(r"Theme (\d)", title)
    for r, row in enumerate(rows, start=1):
        vals = [row[0], row[1],
                f"Theme {theme.group(1)}" if theme else row[2],
                row[3], row[4], row[5], row[6]]
        for c, v in enumerate(vals):
            set_cell(table.cell(r, c), trunc(v, LIT_COLS[c]), 6.5,
                     align=PP_ALIGN.CENTER if c in (0, 2) else PP_ALIGN.LEFT)


# ------------------------------------------------------------------------ slide 5


def _row_groups(s, height_in):
    """Cluster the bar / badge / label / body shapes that form one strip."""
    bars = sorted(
        [sh for sh in s.shapes
         if sh.shape_type == 1 and sh.width == Inches(12.13)
         and abs(sh.height - Inches(height_in)) < Inches(0.01)],
        key=lambda x: x.top)
    groups = []
    for bar in bars:
        grp = [bar]
        for sh in s.shapes:
            if sh is bar or sh.top is None:
                continue
            if bar.top <= sh.top < bar.top + bar.height \
                    and sh.left >= bar.left and sh.width < Inches(12):
                grp.append(sh)
        groups.append(grp)
    return groups


def _retag_badge(grp, label):
    for sh in grp:
        if sh.has_text_frame and abs(sh.width - Inches(0.60)) < Inches(0.01):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = label


def build_gaps(s, C):
    """Template ships 5 strips; the deliverable has 6 gaps, so clone one."""
    groups = _row_groups(s, 1.05)
    groups.append(clone_shapes(s, groups[-1]))         # G6 = copy of the G5 strip

    y0, pitch, h = Inches(1.07), Inches(0.965), Inches(0.86)
    for i, (g, grp) in enumerate(zip(C["gaps"], groups)):
        top = y0 + int(pitch * i)
        for sh in grp:
            sh.top = top + (sh.top - grp[0].top)
            if abs(sh.width - Inches(12.13)) < Inches(0.01):        # bar
                sh.height = h
            elif abs(sh.width - Inches(10.87)) < Inches(0.01):       # body
                sh.height = h
                fill(sh,
                     [[{"t": f"{g['id']} — {g['title']}", "sz": 12.5,
                        "b": True, "c": NAVY, "font": HEAD_FONT}],
                      [{"t": g["gap"], "sz": 8.5, "c": DARK, "line": 0.95}]],
                     margins=(Inches(0.02),) * 4,
                     anchor=MSO_ANCHOR.MIDDLE)
            else:                                                    # badge
                sh.top = top + (h - Inches(0.60)) // 2
        _retag_badge(grp, g["id"])

    fill(by_name(s, "Text 0"),
         [[{"t": "Research Gaps", "sz": 34, "b": True, "c": NAVY,
            "font": HEAD_FONT}]])


# ------------------------------------------------------------------------ slide 6


def build_objs(s, C):
    groups = _row_groups(s, 1.05)
    groups.append(clone_shapes(s, groups[-1]))         # O4 = copy of the O3 strip

    y0, pitch, h = Inches(2.37), Inches(1.12), Inches(1.04)
    for i, (o, grp) in enumerate(zip(C["objs"], groups)):
        top = y0 + int(pitch * i)
        for sh in grp:
            sh.top = top + (sh.top - grp[0].top)
            if abs(sh.width - Inches(12.13)) < Inches(0.01):
                sh.height = h
            elif abs(sh.width - Inches(10.87)) < Inches(0.01):
                sh.height = h
                fill(sh,
                     [[{"t": f"{o['id']} — {o['title']}", "sz": 12.5,
                        "b": True, "c": NAVY, "font": HEAD_FONT}],
                      [{"t": "Target: " + o["target"], "sz": 8.5, "c": DARK,
                        "line": 0.95}],
                      [{"t": f"Addresses {o['gap']}  ·  {o['detail']}",
                        "sz": 8, "i": True, "c": GRAY, "line": 0.95}]],
                     margins=(Inches(0.02),) * 4,
                     anchor=MSO_ANCHOR.MIDDLE)
            elif abs(sh.width - Inches(0.60)) < Inches(0.01):
                sh.top = top + (h - Inches(0.60)) // 2
        _retag_badge(grp, o["id"])

    fill(by_name(s, "Text 5"),
         [[{"t": "MAIN OBJECTIVE", "sz": 15, "b": True, "c": NAVY}],
          [{"t": C["main_obj"], "sz": 11, "c": DARK, "line": 0.95}]],
         anchor=MSO_ANCHOR.MIDDLE)
    fill(by_name(s, "Text 0"),
         [[{"t": "Objectives", "sz": 34, "b": True, "c": NAVY,
            "font": HEAD_FONT}]])


# ------------------------------------------------------------------------ slide 7


def build_refs(s, chunk, part, nparts):
    fill(by_name(s, "Text 0"),
         [[{"t": f"References (APA 7th)  ·  {part}/{nparts}",
            "sz": 34, "b": True, "c": NAVY, "font": HEAD_FONT}]])
    fill(by_name(s, "Text 3"),
         [[{"t": f"{num} {txt}", "sz": 7, "c": DARK, "line": 0.95,
            "space_after": 1.5}] for num, txt in chunk],
         margins=(Inches(0.03),) * 4)


def add_fig_slide(prs, tpl_idx, img, title):
    """Clone a plain slide, drop its body, and centre the figure in the well."""
    s = dup_slide(prs, tpl_idx)
    for nm in ("Shape 2", "Text 3"):                 # background card + body text
        for sh in [x for x in s.shapes if x.name == nm]:
            sh._element.getparent().remove(sh._element)
    fill(by_name(s, "Text 0"),
         [[{"t": title, "sz": 34, "b": True, "c": NAVY, "font": HEAD_FONT}]])

    from PIL import Image
    with Image.open(img) as im:
        ar = im.width / im.height

    well_l, well_t = Inches(0.60), Inches(1.35)
    well_w, well_h = Inches(12.13), Inches(5.40)
    w, h = well_w, int(well_w / ar)
    if h > well_h:
        h, w = well_h, int(well_h * ar)
    s.shapes.add_picture(img, well_l + (well_w - w) // 2,
                         well_t + (well_h - h) // 2, w, h)
    return s


# ------------------------------------------------------------------------- main


def reorder(prs, order):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    pos = {id(sl): i for i, sl in enumerate(prs.slides)}
    keep = [ids[pos[id(sl)]] for sl in order]
    for el in ids:
        lst.remove(el)
    for el in keep:
        lst.append(el)


def number_slides(prs):
    """Sequential page numbers in the bottom-right box (title slide has none)."""
    for n, s in enumerate(prs.slides):
        if n == 0:
            continue
        for sh in s.shapes:
            if sh.has_text_frame and sh.left and sh.top \
                    and abs(sh.left - Inches(11.73)) < Inches(0.05) \
                    and abs(sh.top - Inches(7.00)) < Inches(0.05):
                fill(sh, [[{"t": str(n), "sz": 10, "c": GRAY}]])


def main():
    C = parse_content()
    assert len(C["tech"]) == 6, len(C["tech"])
    assert len(C["gaps"]) == 6, len(C["gaps"])
    assert len(C["objs"]) == 4, len(C["objs"])
    assert len(C["refs"]) == 77, len(C["refs"])
    print("parsed:", {k: len(v) for k, v in C.items() if isinstance(v, list)})

    prs = Presentation(TPL)
    s_title, s_brief, s_tech, s_lit, s_gaps, s_objs, s_refs = list(prs.slides)

    build_title(s_title, C)
    build_brief(s_brief, C)
    build_tech(s_tech, C)
    build_gaps(s_gaps, C)
    build_objs(s_objs, C)

    # ---- literature review: 72 papers, max 11 rows per slide, themes intact
    theme_names = {
        1: "Physical & Acoustic Entropy Sources",
        2: "Entropy Extraction, Conditioning & Statistical Validation",
        3: "Cryptographic Applications & RNG Security",
    }
    lit_slides = []
    for ti, rows in enumerate(C["themes"], start=1):
        chunks = [rows[i:i + 11] for i in range(0, len(rows), 11)]
        for ci, chunk in enumerate(chunks, start=1):
            sl = s_lit if not lit_slides else dup_slide(prs, 3)
            n_all = len(C["themes"][ti - 1])
            build_lit(
                sl, chunk,
                title=f"Literature Review — Theme {ti} ({ci}/{len(chunks)})",
                kicker=f"Theme {ti} · {theme_names[ti]}  —  {n_all} papers",
                note=("Grouped by theme as the template requires. Rows "
                      f"{chunk[0][0].strip('[]')}–{chunk[-1][0].strip('[]')} "
                      f"of {n_all} in this theme · full APA citations on the "
                      "References slides."),
            )
            lit_slides.append(sl)

    # ---- references: 77 entries, 26 per slide
    ref_chunks = [C["refs"][i:i + 26]
                  for i in range(0, len(C["refs"]), 26)]
    ref_slides = [s_refs]
    for _ in ref_chunks[1:]:
        ref_slides.append(dup_slide(prs, 6))
    for i, (sl, chunk) in enumerate(zip(ref_slides, ref_chunks), start=1):
        build_refs(sl, chunk, i, len(ref_chunks))

    # ---- figures (fig5 withheld: it quotes metrics not yet reproduced)
    D = os.path.join(DIAG, "%s.png")
    figs = [
        (1, "fig1_system_architecture", "System Architecture"),
        (2, "fig2_signal_pipeline", "Signal Pipeline"),
        (2, "fig3_entropy_extraction", "Entropy Extraction"),
        (2, "fig4_health_gate", "Source-Health Gate (SP 800-90B style)"),
        (6, "fig6_research_map", "Positioning & Research Map"),
    ]
    fig_slides = [add_fig_slide(prs, i, D % stem, ttl)
                  for i, stem, ttl in figs]

    f1, f2, f3, f4, f6 = fig_slides
    order = [
        s_title, s_brief, f1, s_tech, f2, f3, f4,
        *lit_slides, f6, s_gaps, s_objs, *ref_slides,
    ]
    reorder(prs, order)
    number_slides(prs)

    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", dir=ROOT, delete=False)
    tmp.close()
    prs.save(tmp.name)
    os.replace(tmp.name, TPL)
    print(f"wrote {TPL} — {len(order)} slides")



# ------------------------------------------------------------------------ slide 2


def build_brief(s, C):
    fill(by_name(s, "Text 3"),
         [
             [{"t": "Domain:  ", "sz": 14, "b": True, "c": NAVY},
              {"t": C["domain"], "sz": 11, "c": DARK}],
             [{"t": "Application Area:  ", "sz": 14, "b": True, "c": NAVY},
              {"t": C["app"], "sz": 11, "c": DARK}],
         ])
    for box, idx in (("Text 8", 0), ("Text 13", 1), ("Text 18", 2)):
        txt = C["tri"][idx]
        sz = fit_size(txt, 3.25, 2.70, 11, 7, lines_max=17)
        fill(by_name(s, box),
             [[{"t": txt, "sz": sz, "c": DARK, "line": 1.0}]],
             margins=(Inches(0.02),) * 4)


# ------------------------------------------------------------------------ slide 3


def build_tech(s, C):
    names = ("Text 7", "Text 12", "Text 17", "Text 22", "Text 27", "Text 32")
    bodies = ("Text 8", "Text 13", "Text 18", "Text 23", "Text 28", "Text 33")
    for i, (nm, bd) in enumerate(zip(names, bodies)):
        title, desc = C["tech"][i]
        nsz = fit_size(title, 2.64, 0.60, 17, 11, lines_max=2)
        fill(by_name(s, nm),
             [[{"t": title, "sz": nsz, "b": True, "c": NAVY, "font": HEAD_FONT}]])
        bsz = fit_size(desc, 3.34, 1.05, 13, 8, lines_max=5)
        fill(by_name(s, bd),
             [[{"t": "What it is: ", "sz": bsz, "b": True, "c": NAVY},
               {"t": desc, "sz": bsz, "c": DARK}]])


if __name__ == "__main__":
    main()
